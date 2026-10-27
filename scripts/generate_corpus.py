"""Generate the dummy corpus: corpus/week1 and corpus/week2, plus answer_key.json.

Run from the kb_mvp folder:
    python scripts/generate_corpus.py

Week 2 starts as a byte for byte copy of week 1, then only the changed
documents are regenerated. That way unchanged files keep identical hashes,
which is exactly what change detection needs to see.
"""
import json
import shutil
import sys
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.util import Pt
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

sys.path.insert(0, str(Path(__file__).parent))
from corpus_content import WEEK1, WEEK2_OVERRIDES  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
CORPUS = ROOT / "corpus"


def write_deck(path, slides, font_size=18):
    prs = Presentation()
    layout = prs.slide_layouts[1]  # title and content
    for title, bullets in slides:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for b in bullets[1:]:
            body.add_paragraph().text = b
        for p in body.paragraphs:
            for run in p.runs:
                run.font.size = Pt(font_size)
    prs.save(path)


def write_status(path, title, sections):
    doc = Document()
    doc.add_heading(f"{title}: Weekly Status", level=1)
    for heading, paras in sections:
        doc.add_heading(heading, level=2)
        for text in paras:
            doc.add_paragraph(text)
    doc.save(path)


def write_report(path, title, sections):
    styles = getSampleStyleSheet()
    story = [Paragraph(f"{title}: Phase Report", styles["Title"])]
    for heading, paras in sections:
        story.append(Paragraph(heading, styles["Heading2"]))
        for text in paras:
            story.append(Paragraph(text, styles["BodyText"]))
            story.append(Spacer(1, 6))
    # invariant=1 removes timestamps so identical content gives identical bytes
    SimpleDocTemplate(str(path), pagesize=letter, invariant=1).build(story)


def build_week1():
    week = CORPUS / "week1"
    for slug, p in WEEK1.items():
        folder = week / slug
        folder.mkdir(parents=True, exist_ok=True)
        write_deck(folder / "kickoff_deck.pptx", p["deck"])
        write_status(folder / "weekly_status.docx", p["title"], p["status"])
        write_report(folder / "phase_report.pdf", p["title"], p["report"])


def build_week2():
    src, dst = CORPUS / "week1", CORPUS / "week2"
    shutil.copytree(src, dst)
    for slug, changes in WEEK2_OVERRIDES.items():
        folder = dst / slug
        title = WEEK1[slug]["title"]
        if "status" in changes:
            write_status(folder / "weekly_status.docx", title, changes["status"])
        if changes.get("deck_cosmetic"):
            # same text, bigger font: bytes change, meaning does not
            write_deck(folder / "kickoff_deck.pptx", WEEK1[slug]["deck"], font_size=20)


ANSWER_KEY = {
    "_about": "Ground truth for week 2. Used by the evaluation step.",
    "projects": {
        "medical-records-ocr": {
            "status": "At risk",
            "go_live": "April 2027",
            "decisions": [
                "ClearScan selected as OCR engine",
                "Human review confidence threshold 0.85",
            ],
            "blockers": [
                "Scanner hardware delivery for second clinic",
                "ClearScan API rate limit of 200 forms per hour",
            ],
            "next_steps": [
                "Expand test batch to 500 forms",
                "Draft review queue workflow with clinic staff",
                "Request higher API rate limit from ClearScan",
            ],
        },
        "pharmacy-delivery": {
            "status": "On track",
            "decisions": [
                "MedDash is the courier partner (replaced SwiftRoute)",
                "Delivery window 2 to 6 PM",
            ],
            "blockers": ["Address validation fails for about 6 percent of member records"],
            "next_steps": [
                "Member satisfaction survey in week 3",
                "Clean member address data with data team",
            ],
        },
        "clinical-ai": {
            "status": "On track",
            "decisions": [
                "Note summarization is first pilot workstream",
                "Evaluation rubric approved with five criteria",
            ],
            "blockers": [],
            "next_steps": ["Run first evaluation round on 50 de identified notes"],
        },
    },
    "changes": [
        {"project": "medical-records-ocr", "file": "weekly_status.docx", "meaningful": True,
         "what": "Status moved to At risk and go live moved from March 2027 to April 2027"},
        {"project": "medical-records-ocr", "file": "weekly_status.docx", "meaningful": True,
         "what": "New blocker: ClearScan API rate limit"},
        {"project": "medical-records-ocr", "file": "weekly_status.docx", "meaningful": True,
         "what": "Accuracy improved from 91 to 93 percent"},
        {"project": "medical-records-ocr", "file": "kickoff_deck.pptx", "meaningful": False,
         "what": "Font size change only; text identical"},
        {"project": "pharmacy-delivery", "file": "weekly_status.docx", "meaningful": True,
         "what": "Courier partner changed from SwiftRoute to MedDash"},
        {"project": "pharmacy-delivery", "file": "weekly_status.docx", "meaningful": True,
         "what": "Pilot expanded from two to three clinics; on time rate 93 to 95 percent"},
        {"project": "pharmacy-delivery", "file": "weekly_status.docx", "meaningful": False,
         "what": "Rewording 'will run' to 'plans to run' with no change in meaning"},
        {"project": "clinical-ai", "file": "weekly_status.docx", "meaningful": True,
         "what": "Data access approved; status At risk to On track; blocker resolved"},
        {"project": "clinical-ai", "file": "weekly_status.docx", "meaningful": True,
         "what": "New decision: evaluation rubric approved"},
    ],
    "contradictions": [
        {"project": "medical-records-ocr",
         "what": "Kickoff deck says go live March 2027; weekly status says April 2027",
         "files": ["kickoff_deck.pptx", "weekly_status.docx"]},
        {"project": "pharmacy-delivery",
         "what": "Kickoff deck lists two pilot clinics; weekly status says three",
         "files": ["kickoff_deck.pptx", "weekly_status.docx"]},
    ],
    "unchanged_files": [
        "medical-records-ocr/phase_report.pdf",
        "pharmacy-delivery/kickoff_deck.pptx",
        "pharmacy-delivery/phase_report.pdf",
        "clinical-ai/kickoff_deck.pptx",
        "clinical-ai/phase_report.pdf",
    ],
}


def main():
    if CORPUS.exists():
        shutil.rmtree(CORPUS)
    build_week1()
    build_week2()
    (CORPUS / "answer_key.json").write_text(json.dumps(ANSWER_KEY, indent=2))
    files = sorted(p.relative_to(CORPUS) for p in CORPUS.rglob("*") if p.is_file())
    print(f"Generated {len(files)} files in {CORPUS}")
    for f in files:
        print("  ", f)


if __name__ == "__main__":
    main()
