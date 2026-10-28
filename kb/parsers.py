"""Turn .pptx, .docx and .pdf files into Chunks with exact source locations.

Chunking rule for the MVP:
  pptx -> one chunk per slide (title plus bullets), since a slide is one idea
  docx -> one chunk per paragraph, grouped under the nearest heading
  pdf  -> one chunk per paragraph, headings detected by font size and
          wrapped lines joined back into a single paragraph
"""
from pathlib import Path

import pdfplumber
from docx import Document
from pptx import Presentation

from kb.models import Chunk, normalize


def parse_pptx(path: Path, project: str) -> list[Chunk]:
    chunks = []
    prs = Presentation(path)
    for n, slide in enumerate(prs.slides, start=1):
        title = slide.shapes.title.text if slide.shapes.title else f"Slide {n}"
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape != slide.shapes.title:
                lines += [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
        text = title + ". " + "; ".join(lines) if lines else title
        chunks.append(Chunk(project, path.name, "pptx", title, 1, {"slide": n}, text))
    return chunks


def parse_docx(path: Path, project: str) -> list[Chunk]:
    chunks, section, pos = [], "Body", 0
    for n, para in enumerate(Document(path).paragraphs, start=1):
        text = normalize(para.text)
        if not text:
            continue
        if para.style.name.lower().startswith(("heading", "title")):
            section, pos = text, 0
            continue
        pos += 1
        chunks.append(Chunk(project, path.name, "docx", section, pos, {"paragraph": n}, text))
    return chunks


def parse_pdf(path: Path, project: str) -> list[Chunk]:
    chunks, section, pos = [], "Body", 0
    with pdfplumber.open(path) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            lines = page.extract_text_lines()
            if not lines:
                continue
            body_size = min(round(l["chars"][0]["size"], 1) for l in lines)
            para, prev_top, para_no = [], None, 0

            def flush():
                nonlocal para, pos, para_no
                if para:
                    pos += 1
                    para_no += 1
                    chunks.append(Chunk(project, path.name, "pdf", section, pos,
                                        {"page": page_no, "paragraph": para_no}, " ".join(para)))
                    para = []

            for line in lines:
                size = round(line["chars"][0]["size"], 1)
                if size > body_size:          # bigger font means a heading
                    flush()
                    section, pos, prev_top = line["text"], 0, None
                    continue
                gap = None if prev_top is None else line["top"] - prev_top
                if gap is not None and gap > size * 1.6:   # extra space means a new paragraph
                    flush()
                para.append(line["text"])
                prev_top = line["top"]
            flush()
    return chunks


PARSERS = {".pptx": parse_pptx, ".docx": parse_docx, ".pdf": parse_pdf}


def parse_file(path: Path, project: str) -> list[Chunk]:
    parser = PARSERS.get(path.suffix.lower())
    return parser(path, project) if parser else []
