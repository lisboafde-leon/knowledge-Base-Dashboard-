"""Phase 5b: on demand draft status report for one project, as a Word file.

Usage, from the kb_mvp folder (after python -m kb.publish):
    python -m kb.report --project medical-records-ocr     # build one report now
    python -m kb.report --request medical-records-ocr     # add a row to ReportRequests
    python -m kb.report --process                         # build every Pending request

The request and process pair mirrors the planned SharePoint flow: someone adds a
row to the Report Requests list, the scheduled job picks it up, writes the
draft, and marks the row Ready with a link. Every report is stamped DRAFT and
must be reviewed by a person before it is shared.
"""
import argparse
import getpass
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from kb.publish import OUTPUT, read_table, write_table

REPORTS = OUTPUT / "reports"
STATUS_COLOR = {"On track": RGBColor(0x1E, 0x8E, 0x3E), "At risk": RGBColor(0xC2, 0x7C, 0x0E),
                "Off track": RGBColor(0xC0, 0x39, 0x2B)}


class SourceList:
    """Numbers each distinct source once, like footnotes: [1], [2], ..."""

    def __init__(self):
        self.items = []

    def cite(self, file, location, evidence=""):
        key = (file, location)
        for i, (k, _) in enumerate(self.items, start=1):
            if k == key:
                return i
        self.items.append((key, evidence))
        return len(self.items)


def build_report(project_id: str, requested_by: str = "") -> Path:
    projects = {p["ProjectID"]: p for p in read_table("Projects")}
    if project_id not in projects:
        raise SystemExit(f"Unknown project '{project_id}'. Run python -m kb.publish first.")
    p = projects[project_id]
    facts = [f for f in read_table("Facts") if f["ProjectID"] == project_id]
    changes = [c for c in read_table("Changes") if c["ProjectID"] == project_id and c["Meaningful"] == "True"]
    conflicts = [c for c in read_table("Conflicts") if c["ProjectID"] == project_id]
    sources = SourceList()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Arial", Pt(10.5)

    banner = doc.add_paragraph()
    run = banner.add_run("DRAFT: generated automatically. Review every line against its source before sharing.")
    run.bold, run.font.color.rgb = True, RGBColor(0xC0, 0x39, 0x2B)

    doc.add_heading(f"{p['Title']}: Status Report", level=1)
    meta = doc.add_paragraph()
    meta.add_run(f"Prepared {datetime.now():%B %d, %Y} from {p['Snapshot']} documents by {p['Model']}"
                 + (f" at the request of {requested_by}" if requested_by else "") + ".").italic = True

    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for label, value in [("Status", p["Status"]), ("Target go live", p["GoLive"] or "Not stated"),
                         ("Open blockers", p["OpenBlockers"]), ("Changes this period", p["MeaningfulChanges"]),
                         ("Conflicts to resolve", p["Conflicts"]),
                         ("Unsupported claims", p["UnsupportedClaims"])]:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = str(value)
        if label == "Status" and value in STATUS_COLOR:
            r = cells[1].paragraphs[0].runs[0]
            r.bold, r.font.color.rgb = True, STATUS_COLOR[value]

    def bullet(text, file, location, evidence="", unsupported=False):
        n = sources.cite(file, location, evidence)
        para = doc.add_paragraph(style="List Bullet")
        para.add_run(text)
        para.add_run(f" [{n}]").font.superscript = True
        if unsupported:
            flag = para.add_run("  (unsupported: verify)")
            flag.bold, flag.font.color.rgb = True, RGBColor(0xC0, 0x39, 0x2B)

    purpose = next((f for f in facts if f["Category"] == "Purpose"), None)
    if purpose:
        doc.add_heading("Purpose", level=2)
        bullet(purpose["Text"], purpose["SourceFile"], purpose["SourceLocation"], purpose["Evidence"])

    doc.add_heading("What changed this period", level=2)
    if not changes:
        doc.add_paragraph("No meaningful changes since the last snapshot.")
    for c in changes:
        bullet(c["Summary"], c["SourceFile"], c["SourceLocation"], c["NewText"])

    for heading, cat in [("Decisions", "Decision"), ("Blockers", "Blocker"),
                         ("Next steps", "Next step"), ("Risks", "Risk")]:
        rows = [f for f in facts if f["Category"] == cat]
        doc.add_heading(heading, level=2)
        if not rows:
            doc.add_paragraph("None recorded.")
        for f in rows:
            bullet(f["Text"], f["SourceFile"], f["SourceLocation"], f["Evidence"], f["Supported"] != "True")

    if conflicts:
        doc.add_heading("Conflicts between documents (needs owner decision)", level=2)
        for c in conflicts:
            doc.add_paragraph(f"{c['Text']} (files: {c['SourceFiles']})", style="List Bullet")

    doc.add_heading("Sources", level=2)
    for i, ((file, location), evidence) in enumerate(sources.items, start=1):
        para = doc.add_paragraph()
        para.add_run(f"[{i}] {file}, {location}").bold = True
        if evidence:
            para.add_run(f'\n"{evidence}"').italic = True
    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = "Automatically generated draft. Not for distribution until reviewed."
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    REPORTS.mkdir(parents=True, exist_ok=True)
    path = REPORTS / f"{project_id}_{p['Snapshot']}_{datetime.now():%Y%m%d_%H%M%S}.docx"
    doc.save(path)
    return path


def add_request(project_id: str, requested_by: str = "") -> dict:
    rows = read_table("ReportRequests")
    row = {"RequestID": f"RR{len(rows) + 1:03d}", "ProjectID": project_id,
           "RequestedBy": requested_by or getpass.getuser(),
           "RequestedAt": datetime.now().isoformat(timespec="seconds"),
           "Status": "Pending", "ReportFile": "", "CompletedAt": ""}
    write_table("ReportRequests", rows + [row])
    return row


def process_requests() -> list[dict]:
    rows = read_table("ReportRequests")
    done = []
    for r in rows:
        if r["Status"] != "Pending":
            continue
        try:
            path = build_report(r["ProjectID"], r["RequestedBy"])
            r.update(Status="Ready", ReportFile=str(path.relative_to(OUTPUT.parent)))
        except SystemExit as e:
            r.update(Status=f"Failed: {e}")
        r["CompletedAt"] = datetime.now().isoformat(timespec="seconds")
        done.append(r)
    write_table("ReportRequests", rows)
    return done


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--project", help="build a report for this project now")
    ap.add_argument("--request", help="queue a report request for this project")
    ap.add_argument("--by", default="", help="who is asking (default: your user name)")
    ap.add_argument("--process", action="store_true", help="build every pending request")
    args = ap.parse_args()
    if args.request:
        r = add_request(args.request, args.by)
        print(f"Queued {r['RequestID']} for {r['ProjectID']} (Pending)")
    if args.project:
        print(f"Draft report saved to {build_report(args.project, args.by)}")
    if args.process:
        done = process_requests()
        print(f"Processed {len(done)} request(s)")
        for r in done:
            print(f"   {r['RequestID']} {r['ProjectID']}: {r['Status']} {r['ReportFile']}")
    if not (args.request or args.project or args.process):
        ap.print_help()


if __name__ == "__main__":
    main()
