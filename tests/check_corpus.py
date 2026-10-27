"""Sanity check for Phase 1: confirms the corpus behaves the way the answer key claims.

Run from the kb_mvp folder:
    python tests/check_corpus.py
"""
import hashlib
import json
from pathlib import Path

from docx import Document
from pptx import Presentation

CORPUS = Path(__file__).resolve().parent.parent / "corpus"


def file_hash(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def pptx_text(p):
    prs = Presentation(p)
    return "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)


def docx_text(p):
    return "\n".join(par.text for par in Document(p).paragraphs)


def main():
    key = json.loads((CORPUS / "answer_key.json").read_text())
    w1, w2 = CORPUS / "week1", CORPUS / "week2"
    ok = True

    for rel in key["unchanged_files"]:
        same = file_hash(w1 / rel) == file_hash(w2 / rel)
        print(f"{'PASS' if same else 'FAIL'}  unchanged bytes: {rel}")
        ok &= same

    deck = "medical-records-ocr/kickoff_deck.pptx"
    bytes_differ = file_hash(w1 / deck) != file_hash(w2 / deck)
    text_same = pptx_text(w1 / deck) == pptx_text(w2 / deck)
    print(f"{'PASS' if bytes_differ and text_same else 'FAIL'}  cosmetic deck: bytes differ, text identical")
    ok &= bytes_differ and text_same

    for slug in ["medical-records-ocr", "pharmacy-delivery", "clinical-ai"]:
        rel = f"{slug}/weekly_status.docx"
        changed = docx_text(w1 / rel) != docx_text(w2 / rel)
        print(f"{'PASS' if changed else 'FAIL'}  status text changed: {rel}")
        ok &= changed

    print("\nAll checks passed." if ok else "\nSome checks failed.")


if __name__ == "__main__":
    main()
