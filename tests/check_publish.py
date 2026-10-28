"""Checks for Phase 5: tables, project pages, report builder and dashboard.

Run from the kb_mvp folder:
    python -m kb.run
    python tests/check_publish.py

Uses whatever python -m kb.run produced (the mock model by default).
"""
import csv
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402

from kb.publish import COLUMNS, PAGES, TABLES, read_table  # noqa: E402
from kb.report import add_request, build_report, process_requests  # noqa: E402

key = json.loads((ROOT / "corpus" / "answer_key.json").read_text())["projects"]
results = []


def check(name, cond):
    results.append(bool(cond))
    print(f"{'PASS' if cond else 'FAIL'}  {name}")


# Tables
for name, cols in COLUMNS.items():
    path = TABLES / f"{name}.csv"
    header = next(csv.reader(open(path))) if path.exists() else []
    check(f"{name}.csv exists with SharePoint list columns", header == cols)

projects = {p["ProjectID"]: p for p in read_table("Projects")}
facts, changes, conflicts = read_table("Facts"), read_table("Changes"), read_table("Conflicts")
check("one Projects row per project", set(projects) == set(key))
check("every project status matches the answer key",
      all(projects[p]["Status"] == key[p]["status"] for p in key))
check("OCR go live published as April 2027", projects["medical-records-ocr"]["GoLive"] == "April 2027")
check("blocker counts match the answer key",
      all(int(projects[p]["OpenBlockers"]) == len(key[p]["blockers"]) for p in key))
check("every fact row points to a source file and location",
      all(f["SourceFile"] and f["SourceLocation"] and f["ChunkID"] for f in facts))
check("12 change rows, 11 meaningful", len(changes) == 12 and sum(c["Meaningful"] == "True" for c in changes) == 11)
check("change summaries no longer repeat labels like 'Blocker: Blocker:'",
      not any("Blocker: Blocker" in c["Summary"] or ": Next:" in c["Summary"] for c in changes))
check("conflict row lists both source files",
      any("kickoff_deck.pptx" in c["SourceFiles"] and "weekly_status.docx" in c["SourceFiles"] for c in conflicts))

# Project pages
page = (PAGES / "medical-records-ocr.md").read_text() if (PAGES / "medical-records-ocr.md").exists() else ""
check("project page exists for every project", all((PAGES / f"{p}.md").exists() for p in key))
check("project page shows status, go live, sources and the conflict",
      "At risk" in page and "April 2027" in page and "_Source:" in page and "March 2027" in page)

# Report builder
path = build_report("medical-records-ocr", "Test")
text = "\n".join(p.text for p in Document(path).paragraphs)
check("report is stamped DRAFT", text.startswith("DRAFT"))
check("report includes the go live change and the new blocker", "April 2027" in text and "rate limit" in text)
check("every bullet carries a citation number",
      all("[" in p.text for p in Document(path).paragraphs if p.style.name == "List Bullet"
          and "differs between documents" not in p.text))
n_sources = text.count("weekly_status.docx, paragraph") + text.count("kickoff_deck.pptx, slide") \
    + text.count("phase_report.pdf, page")
check("sources section lists each cited location with its quote", n_sources >= 10 and '"' in text)

# Request flow, the same path the dashboard button uses
before = len(read_table("ReportRequests"))
add_request("pharmacy-delivery", "Test")
check("a report request is queued as Pending",
      read_table("ReportRequests")[-1]["Status"] == "Pending" and len(read_table("ReportRequests")) == before + 1)
done = process_requests()
last = read_table("ReportRequests")[-1]
check("processing builds the file and marks the request Ready",
      done and last["Status"] == "Ready" and (ROOT / last["ReportFile"]).exists())

# Dashboard renders without errors and the button works
try:
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file(str(ROOT / "dashboard" / "app.py"), default_timeout=60).run()
    check("dashboard loads with no exceptions", not at.exception)
    check("dashboard shows the project title", any("Project Knowledge Base" in t.value for t in at.title))
    at.button[0].click().run()
    check("dashboard 'Generate report' button produces a download",
          not at.exception and any("Draft ready" in s.value for s in at.success))
except ImportError:
    print("SKIP  dashboard checks: streamlit is not installed")

print(f"\n{sum(results)}/{len(results)} checks passed.")
